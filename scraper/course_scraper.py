"""
教务系统课表爬虫 v6.0
最终确认的API:
  POST https://jwxt.csmzxy.edu.cn/jwglxt/kbcx/xskbcxMobile_cxXsgrkb.html
  Body: xnm=2025&xqm=12&zs=10&kblx=1&doType=app
  -> 返回 JSON（含 kbList + xsxx）
"""

import json, requests
from datetime import datetime
from pathlib import Path


class CourseScraper:
    """教务系统课表爬虫"""

    API_URL = 'https://jwxt.csmzxy.edu.cn/jwglxt/kbcx/xskbcxMobile_cxXsgrkb.html'
    API_PARAMS = {'xnm': '2025', 'xqm': '12', 'zs': '10', 'kblx': '1', 'doType': 'app'}
    current_week = 10

    def __init__(self):
        self.session = requests.Session()
        self.logged_in = False
        self.student_info = {}
        self.schedule_data = []
        self.debug_info = []
        self.cookie_str = ''

    def set_cookies(self, cookie_str: str) -> bool:
        """设置Cookie并获取课表"""
        self.cookie_str = cookie_str
        self.debug_info = ['设置Cookie...']

        for item in cookie_str.split(';'):
            item = item.strip()
            if '=' in item:
                k, v = item.split('=', 1)
                self.session.cookies.set(k.strip(), v.strip())

        self.debug_info.append(f'已设置 {len(self.session.cookies)} 个cookie')

        headers = {
            'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36 Chrome/129.0.6668.101 wxwork/5.0.8',
            'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
            'X-Requested-With': 'XMLHttpRequest',
            'Origin': 'https://jwxt.csmzxy.edu.cn',
            'Referer': 'https://jwxt.csmzxy.edu.cn/jwglxt/kbcx/xskbcxMobile_cxTimeTableIndex.html?gnmkdm=Y253511',
            'Accept': 'application/json, text/plain, */*',
        }

        try:
            resp = self.session.post(
                self.API_URL,
                data=self.API_PARAMS,
                headers=headers,
                timeout=15,
            )
            if resp.status_code == 200:
                data = resp.json()
                if isinstance(data, dict) and 'kbList' in data:
                    self.logged_in = True
                    self.student_info = data.get('xsxx', {})
                    self.schedule_data = data.get('kbList', [])
                    name = self.student_info.get('XM', '')
                    cls = self.student_info.get('BJMC', '')
                    self.debug_info.append(f'成功! {name} ({cls}) 课表 {len(self.schedule_data)} 条')
                    return True
                else:
                    self.debug_info.append('响应中无kbList字段')
                    self.debug_info.append(f'响应内容: {resp.text[:200]}')
            else:
                self.debug_info.append(f'API返回状态码: {resp.status_code}')
                self.debug_info.append(f'响应: {resp.text[:150]}')
        except Exception as e:
            self.debug_info.append(f'错误: {str(e)[:60]}')

        self.debug_info.append('Cookie无效或API不可访问')
        return False

    def import_from_json(self, json_str: str) -> bool:
        """从JSON字符串导入"""
        self.debug_info = ['从JSON导入...']
        try:
            data = json.loads(json_str)
            if isinstance(data, str):
                data = json.loads(data)
        except json.JSONDecodeError:
            self.debug_info.append('JSON格式错误')
            return False

        kbList = None
        xsxx = {}
        if 'kbList' in data:
            kbList = data['kbList']
            xsxx = data.get('xsxx', {})
        elif isinstance(data, list):
            kbList = data

        if not kbList or not isinstance(kbList, list) or len(kbList) == 0:
            self.debug_info.append('未找到课表数据')
            return False

        self.logged_in = True
        self.schedule_data = kbList
        self.student_info = xsxx
        self.debug_info.append(f'导入成功! 课表 {len(kbList)} 条')
        return True

    def set_week(self, week_num: int) -> bool:
        """设置当前周次并重新获取"""
        self.current_week = week_num
        self.API_PARAMS['zs'] = str(week_num)
        return self.set_cookies(self.cookie_str) if self.cookie_str else False

    def _weeks_match(self, weeks_str: str, target_week: int) -> bool:
        """判断某课程的周次字符串是否包含目标周"""
        if not weeks_str:
            return True
        import re
        for part in weeks_str.split(','):
            part = part.strip()
            is_odd = '单' in part
            is_even = '双' in part
            nums = re.findall(r'(\d+)[-–]*(\d*)', part)
            for n in nums:
                start = int(n[0])
                end_str = n[1]
                end = int(end_str) if end_str else start
                for w in range(start, end + 1):
                    if is_odd and w % 2 == 0:
                        continue
                    if is_even and w % 2 == 1:
                        continue
                    if w == target_week:
                        return True
        return False

    def fetch_schedule(self, week: int = 0) -> list:
        """解析课表（week>0 则按周过滤）"""
        if not self.logged_in:
            raise Exception('未设置数据')

        courses = []
        for item in self.schedule_data:
            try:
                name = item.get('kcmc', '').strip()
                teacher = item.get('xm', '').strip()
                loc = f"{item.get('lh', '')} {item.get('cdmc', '')}".strip()
                day = int(item.get('xqj', 0))
                period_str = str(item.get('jcs', '') or item.get('jcor', '')).strip()
                weeks_str = item.get('zcd', '').strip()
                nature = item.get('kcxz', '')

                if not name:
                    continue
                if week > 0 and not self._weeks_match(weeks_str, week):
                    continue

                start_p, end_p = 1, 1
                if period_str:
                    p = period_str.replace('节', '')
                    if '上午' in p:
                        start_p, end_p = 1, 4
                    elif '下午' in p:
                        start_p, end_p = 5, 8
                    elif '晚上' in p:
                        start_p, end_p = 9, 11
                    elif '-' in p:
                        parts = p.split('-')
                        try:
                            start_p, end_p = int(parts[0]), int(parts[1])
                        except ValueError:
                            pass
                    else:
                        try:
                            start_p = end_p = int(p)
                        except ValueError:
                            start_p = end_p = 12

                for p in range(start_p, end_p + 1):
                    courses.append({
                        'name': name, 'teacher': teacher, 'location': loc,
                        'weekday': day, 'period': p, 'weeks': weeks_str, 'nature': nature,
                    })
            except:
                continue

        seen = set()
        unique = []
        for c in courses:
            key = (c['name'], c['weekday'], c['period'], c['teacher'])
            if key not in seen:
                seen.add(key)
                unique.append(c)
        return unique

    def fetch_semester(self) -> list:
        """获取整学期课表（遍历1-20周API去重，返回统一格式）"""
        if not self.logged_in or not self.cookie_str:
            return self.fetch_schedule()

        self.debug_info = ['获取整学期课表...']
        original_week = self.current_week
        all_raw_items = []

        for week in range(1, 21):
            self.API_PARAMS['zs'] = str(week)
            try:
                resp = self.session.post(
                    self.API_URL,
                    data=self.API_PARAMS,
                    headers={
                        'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; WOW64) AppleWebKit/537.36',
                        'Content-Type': 'application/x-www-form-urlencoded; charset=UTF-8',
                        'X-Requested-With': 'XMLHttpRequest',
                        'Referer': 'https://jwxt.csmzxy.edu.cn/jwglxt/kbcx/xskbcxMobile_cxTimeTableIndex.html',
                    },
                    timeout=10,
                )
                if resp.status_code == 200:
                    data = resp.json()
                    items = data.get('kbList', [])
                    if items:
                        all_raw_items.extend(items)
                    self.debug_info.append(f'  第{week}周: {len(items)} 门')
                else:
                    self.debug_info.append(f'  第{week}周: HTTP {resp.status_code}')
                    self.debug_info.append(f'    响应: {resp.text[:150]}')
            except Exception as e:
                self.debug_info.append(f'  第{week}周: {str(e)[:40]}')

        self.API_PARAMS['zs'] = str(original_week)
        self.current_week = original_week

        seen = set()
        unique = []
        for item in all_raw_items:
            key = (item.get('kcmc', ''), item.get('xqj', ''),
                   str(item.get('jcs', '')), item.get('xm', ''))
            if key not in seen:
                seen.add(key)
                unique.append(item)

        self.schedule_data = unique
        courses = self.fetch_schedule()
        self.debug_info.append(f'合并去重完成: {len(courses)} 门课程')
        return courses

    def get_student_info(self) -> dict:
        return self.student_info

    def get_debug_info(self) -> list:
        return self.debug_info

    def export_to_excel(self, courses: list, output_path: str = '') -> str:
        import openpyxl
        from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = '课程表'

        hf = PatternFill(start_color='2B6CB0', end_color='2B6CB0', fill_type='solid')
        hfont = Font(bold=True, color='FFFFFF', size=11, name='微软雅黑')
        tb = Border(left=Side(style='thin', color='CCCCCC'), right=Side(style='thin', color='CCCCCC'),
                     top=Side(style='thin', color='CCCCCC'), bottom=Side(style='thin', color='CCCCCC'))
        ca = Alignment(horizontal='center', vertical='center', wrap_text=True)

        if self.student_info:
            ws.cell(row=1, column=1, value=f"学生: {self.student_info.get('XM', '')}")
            ws.cell(row=1, column=3, value=f"班级: {self.student_info.get('BJMC', '')}")
            for c in range(1, 9):
                ws.cell(row=1, column=c).font = Font(bold=True, size=10, color='2B6CB0')

        days = ['', '周一', '周二', '周三', '周四', '周五', '周六', '周日']
        for col, day in enumerate(days, 1):
            cell = ws.cell(row=2, column=col, value=day)
            cell.font = hfont; cell.fill = hf; cell.alignment = ca; cell.border = tb

        periods = ['第1节\n08:00', '第2节\n08:55', '第3节\n10:00', '第4节\n10:55',
                    '第5节\n14:00', '第6节\n14:55', '第7节\n16:00', '第8节\n16:55',
                    '第9节\n19:00', '第10节\n19:55']
        for row, p in enumerate(periods, 3):
            cell = ws.cell(row=row, column=1, value=p)
            cell.font = Font(color='667788', size=10); cell.alignment = ca; cell.border = tb

        colors = {'专必': 'D6EAF8', '公基': 'E8F8F5', '公选': 'FEF9E7', '专任': 'F4ECF7'}
        for course in courses:
            wd, pr = course.get('weekday', 0), course.get('period', 0)
            if 1 <= wd <= 7 and 1 <= pr <= 10:
                parts = [course['name']]
                if course.get('teacher'): parts.append(course['teacher'])
                if course.get('location'): parts.append(course['location'])
                if course.get('weeks'): parts.append(course['weeks'])
                cell = ws.cell(row=pr+2, column=wd+1, value='\n'.join(parts))
                cell.alignment = ca; cell.border = tb
                color = colors.get(course.get('nature', ''), 'EBF5FB')
                cell.fill = PatternFill(start_color=color, end_color=color, fill_type='solid')

        ws.column_dimensions['A'].width = 14
        for c in 'BCDEFGH': ws.column_dimensions[c].width = 18
        for r in range(3, 13): ws.row_dimensions[r].height = 60

        if not output_path:
            output_path = f'课程表_{datetime.now():%Y%m%d_%H%M%S}.xlsx'
        wb.save(output_path)
        return output_path

    def check_connection(self, url: str) -> dict:
        result = {'connected': False, 'system': '', 'error': ''}
        try:
            r = self.session.get('https://jwxt.csmzxy.edu.cn', timeout=10)
            result['connected'] = r.status_code in (200, 302)
            result['system'] = '长沙民政职院教务系统'
        except Exception as e:
            result['error'] = str(e)
        return result
